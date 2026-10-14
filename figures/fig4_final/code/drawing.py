"""Editable PowerPoint primitives in inches; no rasterized mechanism panels."""

from pathlib import Path
import math
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont

ROOT = Path(__file__).resolve().parent.parent
FONTDIR = ROOT / "assets" / "liberation-fonts-ttf-2.1.5"
C = dict(
    ink="20262E",
    muted="68727D",
    blue="4C78A8",
    redox="2C7FB8",
    energy="E69F00",
    orange="D55E00",
    stop="C43C39",
    green="009E73",
    rule="D7DCE2",
    purple="7B6FD0",
    pale_blue="EDF4FA",
    pale_yellow="FFF5D9",
    pale_orange="FCF0E8",
    pale_green="EDF6F0",
    white="FFFFFF",
)


def rgb(c):
    return RGBColor.from_string(C.get(c, c))


def E(v):
    return int(round(v * 914400))


class Drawing:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(16)
        self.prs.slide_height = Inches(32 / 3)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb("white")
        self.texts = []

    def group(self, name):
        g = self.slide.shapes.add_group_shape()
        g.name = name
        return g.shapes

    def shift_group(self, name, dy):
        """Rigid reflow: preserve circle aspect ratios, fonts and arrow geometry."""
        group = next(s for s in self.slide.shapes if s.name == name)
        group.top += E(dy)
        for item in self.texts:
            if item.get("owner") == name:
                item["y"] += dy

    def style(self, s, color="blue", width=1, fill=None):
        # Neutralize theme effects: scientific figures must have flat strokes, no shadows.
        for child in list(s._element):
            if child.tag.endswith("}style"):
                s._element.remove(child)
        effects = OxmlElement("a:effectLst")
        s._element.spPr.append(effects)
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if color:
            s.line.color.rgb = rgb(color)
            s.line.width = Pt(width)
        else:
            s.line.fill.background()
        return s

    def text(
        self,
        parent,
        x,
        y,
        w,
        h,
        text,
        size=12,
        color="ink",
        bold=False,
        italic=False,
        align="center",
        name=None,
    ):
        if re.search(r"[⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉]", text):
            return self.chemical_text(
                parent, x, y, w, h, text, size, color, bold, italic, align, name
            )
        s = parent.add_textbox(E(x), E(y), E(w), E(h))
        s.name = name or text.replace("\n", " ")[:65]
        tf = s.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {
                "center": PP_ALIGN.CENTER,
                "left": PP_ALIGN.LEFT,
                "right": PP_ALIGN.RIGHT,
            }[align]
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.0
            supers = "⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹"
            subs = "₀₁₂₃₄₅₆₇₈₉"
            mapping = str.maketrans(supers + subs, "+-0123456789" + "0123456789")
            for token in re.findall(
                r"[⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+|[₀₁₂₃₄₅₆₇₈₉]+|[^⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉]+", line
            ):
                r = p.add_run()
                r.text = token.translate(mapping)
                r.font.name = "Liberation Sans"
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.italic = italic
                r.font.color.rgb = rgb(color)
                if token[0] in supers + subs:
                    r.font.size = Pt(size * 0.8)
                    r._r.get_or_add_rPr().set(
                        "baseline", "30000" if token[0] in supers else "-18000"
                    )
        self.texts.append(
            dict(
                text=text,
                x=x,
                y=y,
                w=w,
                h=h,
                size=size,
                bold=bold,
                italic=italic,
                owner=getattr(parent._parent, "name", "slide"),
            )
        )
        return s

    def chemical_text(
        self, parent, x, y, w, h, text, size, color, bold, italic, align, name
    ):
        """Position editable ASCII superscripts explicitly, avoiding missing glyphs
        and Office's additional automatic shrinkage of baseline-shifted runs.
        """
        group = parent.add_group_shape()
        group.name = name or text.replace("\n", " ")[:65]
        supers = "⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹"
        subs = "₀₁₂₃₄₅₆₇₈₉"
        mapping = str.maketrans(supers + subs, "+-0123456789" + "0123456789")
        variant = (
            "BoldItalic"
            if bold and italic
            else "Bold" if bold else "Italic" if italic else "Regular"
        )
        fonts = {
            factor: ImageFont.truetype(
                str(FONTDIR / f"LiberationSans-{variant}.ttf"),
                round(size * factor * 20),
            )
            for factor in [1, 0.72]
        }
        lines = text.split("\n")
        line_h = size * 1.08 / 72
        widths = []
        for i, line in enumerate(lines):
            pieces = []
            for token in re.findall(
                r"[⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+|[₀₁₂₃₄₅₆₇₈₉]+|[^⁺⁻⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉]+", line
            ):
                factor = 0.72 if token[0] in supers + subs else 1
                shift = (
                    -0.20 * size / 72
                    if token[0] in supers
                    else 0.28 * size / 72 if token[0] in subs else 0
                )
                raw = token.translate(mapping)
                width = fonts[factor].getlength(raw) / 20 / 72
                pieces.append((raw, factor, shift, width))
            total = sum(p[3] for p in pieces)
            widths.append(total * 72)
            xx = (
                x
                if align == "left"
                else x + w - total if align == "right" else x + (w - total) / 2
            )
            yy = y + (h - len(lines) * line_h) / 2 + i * line_h
            for raw, factor, shift, width in pieces:
                s = group.shapes.add_textbox(
                    E(xx), E(yy + shift), E(width + 0.02), E(line_h)
                )
                s.name = "Chemical text: " + raw
                tf = s.text_frame
                tf.clear()
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = 1.0
                r = p.add_run()
                r.text = raw
                r.font.name = "Liberation Sans"
                r.font.size = Pt(size * factor)
                r.font.bold = bold
                r.font.italic = italic
                r.font.color.rgb = rgb(color)
                xx += width
        self.texts.append(
            dict(
                text=text,
                x=x,
                y=y,
                w=w,
                h=h,
                size=size,
                bold=bold,
                italic=italic,
                measured_run_width_pt=max(widths),
                owner=getattr(parent._parent, "name", "slide"),
            )
        )
        return group

    def shape(
        self, parent, kind, x, y, w, h, color="blue", width=1, fill=None, name=None
    ):
        s = parent.add_shape(kind, E(x), E(y), E(w), E(h))
        self.style(s, color, width, fill)
        if name:
            s.name = name
        return s

    def circle(self, parent, x, y, r=0.052, color="blue", fill="white", width=1.15):
        return self.shape(
            parent, MSO_SHAPE.OVAL, x - r, y - r, 2 * r, 2 * r, color, width, fill
        )

    def rect(
        self,
        parent,
        x,
        y,
        w,
        h,
        color=None,
        width=0.7,
        fill="pale_blue",
        rounding=False,
    ):
        return self.shape(
            parent,
            MSO_SHAPE.ROUNDED_RECTANGLE if rounding else MSO_SHAPE.RECTANGLE,
            x,
            y,
            w,
            h,
            color,
            width,
            fill,
        )

    def path(
        self,
        parent,
        commands,
        color="blue",
        width=2,
        arrow=True,
        dash=False,
        name="Metabolic route",
    ):
        pts = []
        for cmd in commands:
            vals = cmd[1:]
            pts.extend(zip(vals[::2], vals[1::2]))
        xs, ys = zip(*pts)
        x0, y0 = min(xs), min(ys)
        x1, y1 = max(xs), max(ys)
        # Native custom geometry with Bezier segments preserves editability.
        s = parent.add_shape(
            MSO_SHAPE.RECTANGLE, E(x0), E(y0), max(E(x1 - x0), 1), max(E(y1 - y0), 1)
        )
        s.name = name
        sp = s._element.spPr
        for child in list(sp):
            if child.tag.endswith("prstGeom"):
                sp.remove(child)
        geom = OxmlElement("a:custGeom")
        for tag in ["avLst", "gdLst", "ahLst", "cxnLst"]:
            geom.append(OxmlElement("a:" + tag))
        rect = OxmlElement("a:rect")
        for k, v in dict(l="0", t="0", r="r", b="b").items():
            rect.set(k, v)
        geom.append(rect)
        pl = OxmlElement("a:pathLst")
        p = OxmlElement("a:path")
        p.set("w", str(max(E(x1 - x0), 1)))
        p.set("h", str(max(E(y1 - y0), 1)))
        p.set("fill", "none")
        for cmd in commands:
            tag = {"M": "moveTo", "L": "lnTo", "C": "cubicBezTo"}[cmd[0]]
            elem = OxmlElement("a:" + tag)
            for xx, yy in zip(cmd[1::2], cmd[2::2]):
                pt = OxmlElement("a:pt")
                pt.set("x", str(E(xx - x0)))
                pt.set("y", str(E(yy - y0)))
                elem.append(pt)
            p.append(elem)
        pl.append(p)
        geom.append(pl)
        sp.insert(1, geom)
        self.style(s, color, width, None)
        ln = sp.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
        ln.set("cap", "rnd")
        if dash:
            d = OxmlElement("a:prstDash")
            d.set("val", "dash")
            ln.append(d)
        if arrow:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            end.set("w", "med")
            end.set("len", "med")
            ln.append(end)
        return s

    def line(self, parent, x1, y1, x2, y2, **kwargs):
        return self.path(parent, [("M", x1, y1), ("L", x2, y2)], **kwargs)

    def cross(self, parent, x, y, r=0.075):
        self.line(
            parent,
            x - r,
            y - r,
            x + r,
            y + r,
            color="stop",
            width=2.6,
            arrow=False,
            name="MQC stop",
        )
        self.line(
            parent,
            x - r,
            y + r,
            x + r,
            y - r,
            color="stop",
            width=2.6,
            arrow=False,
            name="MQC stop",
        )

    def membrane(self, parent, x, y, w, h=0.36, gap=None, density=0.12):
        n = round(w / density) + 1
        for i in range(n):
            xx = x + w * i / (n - 1)
            if gap and gap[0] - 0.03 < xx < gap[1] + 0.03:
                continue
            r = 0.034
            for yy, sign in [(y, 1), (y + h, -1)]:
                for dx in [-0.018, 0.018]:
                    self.path(
                        parent,
                        [
                            ("M", xx + dx, yy + sign * r),
                            ("L", xx + dx, yy + sign * h * 0.30),
                            ("L", xx + dx + 0.013, yy + sign * h * 0.45),
                        ],
                        width=0.7,
                        arrow=False,
                        color="muted",
                        name="Phospholipid tail",
                    )
                self.circle(parent, xx, yy, r, fill="pale_blue", width=0.8)

    def check_text(self):
        issues = []
        for t in self.texts:
            variant = (
                "BoldItalic"
                if t["bold"] and t["italic"]
                else "Bold" if t["bold"] else "Italic" if t["italic"] else "Regular"
            )
            font = ImageFont.truetype(
                str(FONTDIR / f"LiberationSans-{variant}.ttf"), round(t["size"] * 10)
            )
            wide = t.get(
                "measured_run_width_pt",
                max(font.getlength(s) / 10 for s in t["text"].split("\n")),
            )
            if wide > t["w"] * 72 + 1:
                issues.append(dict(**t, measured_width_pt=round(wide, 2)))
            if len(t["text"].split("\n")) * t["size"] * 1.03 > t["h"] * 72 + 2:
                issues.append(dict(**t, height_warning=True))
        return issues
