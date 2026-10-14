"""Read-only scientific/input checks; write audit artifacts, never optimize GEMs."""

from pathlib import Path
import hashlib
import json
import math
import re
import zipfile
from collections import defaultdict
import pandas as pd
from pptx import Presentation
from cobra.io import read_sbml_model

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "source_data" / "fig4"
AUDIT = ROOT / "review"
AUDIT.mkdir(exist_ok=True)
STEM = "Fig4_final"


def table(name):
    return pd.read_csv(DATA / name)


def close(a, b):
    assert math.isclose(float(a), float(b), rel_tol=0, abs_tol=1e-8), (a, b)


def recurse(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from recurse(shape.shapes)


def clean(eq):
    return re.sub(r"\s+", "", re.sub(r"(\d+)\.0\b", r"\1", eq))


def main():
    checked = []
    manifest = json.loads((ROOT / "source_data" / "source_manifest.json").read_text())
    for item in manifest:
        path = ROOT / item["packaged_path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"], str(
            path
        )
    checked.append(f"{len(manifest)} archived input files: SHA-256 matched")
    diag = table("diagnostic_quantitative_data.csv").set_index("target")
    no = pd.read_csv(
        ROOT / "Fig4_panel_c" / "data" / "no_carbon_objectives.csv"
    ).set_index("objective")
    mapping = {
        "NADH": "NADH",
        "ATP/CTP/GTP/UTP/ITP": "ATP",
        "Urea": "urea_c",
        "Allophanate": "allphn_c",
    }
    assert set(no.index) == set(mapping)
    for label, target in mapping.items():
        close(no.loc[label, "initial"], diag.loc[target, "initial"])
        close(no.loc[label, "post_qc"], diag.loc[target, "final"])
    for t in ["ATP", "CTP", "GTP", "UTP", "ITP"]:
        close(diag.loc[t, "initial"], 96.250)
        close(diag.loc[t, "final"], 0)
    matched = table("matched_glucose_comparison.csv")
    shown = pd.read_csv(ROOT / "Fig4_panel_c" / "data" / "matched_glucose_retest.csv")
    assert len(matched) == len(shown) == 3
    assert list(matched["metric"]) == ["NADH capacity", "ATP capacity", "Biomass"]
    assert list(shown["metric"]) == [
        "Maximum NADH production",
        "Maximum ATP production",
        "Maximum biomass production",
    ]
    for (_, a), (_, b) in zip(matched.iterrows(), shown.iterrows()):
        close(a.pre_mqc, b.initial)
        close(a.post_mqc, b.post_qc)
        close(a.uptake_bound, -b.uptake)
        assert a.carbon_source == b.carbon_source == "D-glucose"
        assert a.other_carbon_sources == b.other_carbon_sources == "closed"
        assert a.unit == b.unit
    points = pd.read_csv(
        ROOT / "Fig4_panel_c" / "output" / "Fig4_panel_c_coordinate_audit.csv"
    )
    assert len(points) == 14
    for _, r in points.iterrows():
        close(r.value, r.artist_x)
        assert 0 <= r.x_axis_fraction <= 1 and 0 <= r.y_axis_fraction <= 1
    checked.append(
        "Panel c: 14 plotted points, source mappings, axis bounds and matched-carbon conditions passed"
    )

    # Reading stoichiometry and formula annotations does not call optimize/pFBA.
    model = read_sbml_model(str(ROOT / "source_data" / "iCN718.xml"))
    source_eq = table("reaction_stoichiometry.csv")
    for _, r in source_eq.iterrows():
        assert clean(r.equation) == clean(
            model.reactions.get_by_id(r.reaction_id).reaction
        ), r.reaction_id
    checked.append(
        "All 12 archived reaction equations match packaged local SBML stoichiometry"
    )

    def coeff(r):
        return {
            m.id: float(v) for m, v in model.reactions.get_by_id(r).metabolites.items()
        }

    assert coeff("RDXR")["nadh_c"] == -1 and "nadph_c" not in coeff("RDXR")
    assert coeff("OMMBLHX3")["2ommbl_c"] == -1 and coeff("OMMBLHX3")["2omhmbl_c"] == 1
    assert coeff("MADH")["co2_c"] == 1 and coeff("MADH")["mmal_c"] == 1
    assert coeff("UREA")["co2_c"] == 1 and coeff("UREA")["nh4_c"] == 2
    assert coeff("ALPHNH")["co2_c"] == 2 and coeff("ALPHNH")["nh4_c"] == 2
    routes = table("causal_path_fluxes.csv")
    net = defaultdict(float)
    for _, r in routes[routes.path == "FDR"].iterrows():
        if r.reaction_id == "ADD_NADH":
            continue
        for met, c in coeff(r.reaction_id).items():
            net[met] += r.flux * c
    net = {k: v for k, v in net.items() if abs(v) > 1e-7}
    assert net == {"nadh_c": 500.0, "h_c": -500.0, "nad_c": -500.0}, net
    checked.append(
        "RDXR cofactor, OMMBLHX3 orientation, MADH CO2 branch and flux-weighted FDR net reaction passed"
    )
    eqs = []
    for rid in list(source_eq.reaction_id) + ["ALPHNH"]:
        r = model.reactions.get_by_id(rid)
        eqs.append(
            {
                "reaction_id": rid,
                "equation": r.reaction,
                "lower_bound": r.lower_bound,
                "upper_bound": r.upper_bound,
                "coefficients_json": json.dumps(coeff(rid), sort_keys=True),
            }
        )
    pd.DataFrame(eqs).to_csv(AUDIT / "local_model_reaction_equations.csv", index=False)
    comparisons = []
    for _, r in table("reaction_revisions.csv").iterrows():
        if r.reaction_id == "EX_h_e":
            continue
        balance = model.reactions.get_by_id(r.reaction_id).check_mass_balance()
        comparisons.append(
            {
                "reaction_id": r.reaction_id,
                "archived_evidence": r.evidence,
                "current_local_model_mass_balance": json.dumps(balance, sort_keys=True),
                "interpretation": "Historical diagnostic annotation is not reproduced by this local model version; specific residuals omitted from final artwork.",
            }
        )
    pd.DataFrame(comparisons).to_csv(
        AUDIT / "historical_vs_local_annotation.csv", index=False
    )

    prs = Presentation(ROOT / (STEM + ".pptx"))
    assert len(prs.slides) == 1
    shapes = list(recurse(prs.slides[0].shapes))
    runs = [
        r
        for s in shapes
        if s.has_text_frame
        for p in s.text_frame.paragraphs
        for r in p.runs
        if r.text
    ]
    assert all(r.font.name == "Liberation Sans" for r in runs)
    text = "\n".join(s.text for s in shapes if s.has_text_frame)
    assert "Boundary normalization" not in text and "Functional re-testing" not in text
    assert "Reducing equivalents" in text and "NADHs" not in text
    for label in [
        "a",
        "b",
        "c",
        "665.802",
        "96.250 each",
        "5.000",
        "NADH output: 1.667",
        "ATPS4r flux: 2.500",
        "Mass / charge imbalance",
    ]:
        assert label in text, label
    assert "charge −2; C +1; O +3" not in text
    assert "H −2" not in text
    arrow = next(s for s in shapes if s.name == "OMMBLHX3:2ommbl_c>2omhmbl_c")
    pts = arrow._element.xpath(".//a:pathLst/a:path/*/a:pt")
    assert int(pts[0].get("x")) > int(pts[-1].get("x"))
    assert arrow._element.xpath('.//a:tailEnd[@type="triangle"]')
    branch = next(
        s for s in shapes if s.name == "ALPHNH:model-derived separate-objective branch"
    )
    assert branch._element.xpath('.//a:prstDash[@val="dash"]')
    with zipfile.ZipFile(ROOT / (STEM + ".pptx")) as z:
        assert (
            len(
                [
                    n
                    for n in z.namelist()
                    if n.startswith("ppt/fonts/") and n.endswith(".fntdata")
                ]
            )
            == 4
        )
    sizes = json.loads((ROOT / "code" / "text_inventory.json").read_text())
    assert min(t["size"] for t in sizes) >= 13.5
    assert json.loads((ROOT / "code" / "text_fit_check.json").read_text()) == []
    checked.append(
        "PPT: OMMBLHX3 arrow direction, model-derived ALPHNH dashed branch, four embedded font faces, minimum base font 13.5 pt and text dimensions passed"
    )
    result = {
        "automated_checks": "PASS",
        "no_model_optimization_run": True,
        "checks": checked,
        "scientific_scope_note": "Historical imbalance annotations differ from the packaged local SBML; final artwork omits exact residuals. Allophanate is a separate diagnostic objective, not a simultaneously traced flux of the urea solution.",
        "visual_review": "Requires inspection of the final rendered PDF and 183 mm proof; see README.md.",
    }
    (AUDIT / "evidence_check.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
